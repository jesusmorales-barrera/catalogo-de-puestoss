from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import argparse
import json
import re

BURGUNDY = "691A32"
GOLD = "BC955B"
WHITE = "FFFFFF"
BLACK = "000000"
LIGHT = "F8F4EE"

# Los encabezados reales de tu archivo están en la fila 3.
HEADER_ROW = 3
DEFAULT_SHEET = "ESTRUCTURA (2)"


def normalize(text):
    if text is None:
        return ""
    text = str(text).replace("\n", " ").strip().upper()
    repl = str.maketrans("ÁÉÍÓÚÜÑ", "AEIOUUN")
    text = text.translate(repl)
    return re.sub(r"\s+", " ", text)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GOLD, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        el = tc_borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def remove_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tc_borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            tc_borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_run(p, text, size=9.5, bold=False, color=BLACK):
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def style_cell(cell, label=False):
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if label:
        set_cell_shading(cell, LIGHT)


def fill_cell(cell, text, bold=False, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_run(p, text if text else " ", size=size, bold=bold)


def set_width(cell, inches):
    cell.width = Inches(inches)


def set_table_widths(table, widths):
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
        for cell in table.columns[index].cells:
            set_width(cell, width)


def add_table(container, rows, cols):
    if container.__class__.__name__ == "_Header":
        return container.add_table(rows=rows, cols=cols, width=Inches(7.4))
    return container.add_table(rows=rows, cols=cols)


def add_section_bar(container, title):
    t = add_table(container, rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    set_cell_shading(cell, BURGUNDY)
    set_cell_margins(cell, top=70, bottom=70, start=120, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, title, size=11.5, bold=True, color=WHITE)


def add_header(container, logo_path):
    # Logotipo institucional horizontal completo, sin recortes ni texto duplicado.
    logo_p = container.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_before = Pt(0)
    logo_p.paragraph_format.space_after = Pt(3)
    if logo_path.exists():
        logo_p.add_run().add_picture(str(logo_path), width=Inches(5.9))

    # Franja institucional completamente guinda.
    line = add_table(container, rows=1, cols=1)
    line.alignment = WD_TABLE_ALIGNMENT.CENTER
    line.autofit = False
    bar = line.cell(0, 0)
    set_width(bar, 7.4)
    set_cell_shading(bar, BURGUNDY)
    remove_cell_borders(bar)
    bar.text = ""
    set_cell_margins(bar, top=15, bottom=15, start=0, end=0)

    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "CATÁLOGO SECTORIAL DE PUESTOS", size=15, bold=True, color=BURGUNDY)

def add_four_column_row(container, values):
    t = add_table(container, rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    widths = [1.25, 3.05, 1.75, 1.35]
    set_table_widths(t, widths)
    row = t.rows[0]

    for i, value in enumerate(values):
        label = i in (0, 2)
        style_cell(row.cells[i], label=label)
        fill_cell(row.cells[i], value, bold=label, size=8.8 if label else 9.2)


def add_two_column_row(container, label, value):
    t = add_table(container, rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    label_cell, value_cell = t.rows[0].cells
    set_table_widths(t, [1.25, 6.15])
    style_cell(label_cell, label=True)
    style_cell(value_cell, label=False)
    fill_cell(label_cell, label, bold=True, size=8.2)
    fill_cell(value_cell, value, size=9.2)


def add_position_data(container, record):
    add_section_bar(container, "DATOS DEL PUESTO")
    add_four_column_row(
        container,
        ("PUESTO", record["CARGO EN ESTRUCTURA"], "NIVEL", record["NIVEL"]),
    )
    add_two_column_row(
        container,
        "DENOMINACIÓN GENÉRICA DEL PUESTO",
        record["DENOMINACION GENERICA DEL PUESTO"],
    )


def add_requirements(doc, record):
    add_section_bar(doc, "REQUISITOS")
    add_four_column_row(
        doc,
        (
            "ESCOLARIDAD REQUERIDA",
            record["ESCOLARIDAD REQUERIDA"],
            "EXPERIENCIA (AÑOS)",
            record["EXPERIENCIA (AÑOS)"],
        ),
    )


def set_row_min_height(row, twips=1100):
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(twips))
    trHeight.set(qn("w:hRule"), "atLeast")
    trPr.append(trHeight)


def add_long_section(doc, title, label, text):
    title = title.replace("\ufffdN", "\u00d3N").replace("GEN\ufffdRICA", "GEN\u00c9RICA")
    label = label.replace("\ufffdN", "\u00d3N")
    add_section_bar(doc, title)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    label_cell, value_cell = t.rows[0].cells
    set_row_min_height(t.rows[0], 1050)
    set_table_widths(t, [1.25, 6.15])
    style_cell(label_cell, label=True)
    style_cell(value_cell, label=False)
    fill_cell(label_cell, label, bold=True, size=9)
    value_cell.text = ""
    paragraphs = split_description_paragraphs(text)
    for index, paragraph_text in enumerate(paragraphs):
        p = value_cell.paragraphs[0] if index == 0 else value_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4 if index < len(paragraphs) - 1 else 1)
        p.paragraph_format.line_spacing = 1.05
        add_run(p, paragraph_text, size=9.2)


def split_description_paragraphs(text):
    """Convierte saltos de línea e incisos en párrafos Word reales.

    Las líneas que continúan una oración se unen con un espacio. Un inciso nuevo,
    una viñeta o una línea que termina en punto o punto y coma cierra el párrafo.
    Así, Word no estira la última línea al justificar el texto.
    """
    raw = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    inline_roman_start = re.compile(r"(?<=\S)\s+(?=[IVXLCDM]+\.\s)")
    raw = inline_roman_start.sub("\n", raw)
    lines = [re.sub(r"\s+", " ", line).strip() for line in raw.split("\n")]
    item_start = re.compile(r"^(?:[•●▪◦�-]|(?:[IVXLCDM]+|\d+)[.)])\s*")
    paragraphs = []
    current = []

    def flush():
        if current:
            paragraphs.append(" ".join(current).strip())
            current.clear()

    for line in lines:
        if not line:
            flush()
            continue
        if item_start.match(line) and current:
            flush()
        current.append(line)
        if line.endswith((".", ";")):
            flush()

    flush()
    return paragraphs or [" "]


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(paragraph, "Página ", size=8, color="666666")
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def setup_section(sec):
    sec.top_margin = Inches(3.05)
    sec.bottom_margin = Inches(0.42)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    sec.header_distance = Inches(0.15)
    sec.footer_distance = Inches(0.18)


def set_record_header(sec, record, logo_path):
    setup_section(sec)
    sec.header.is_linked_to_previous = False
    header = sec.header
    for child in list(header._element):
        header._element.remove(child)
    add_header(header, logo_path)
    add_position_data(header, record)


def setup_doc(doc):
    sec = doc.sections[0]
    setup_section(sec)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9.5)
    fp = sec.footer.paragraphs[0]
    add_page_number(fp)


def load_records_excel(excel_path, sheet_name=DEFAULT_SHEET):
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"No existe la hoja '{sheet_name}'. Hojas disponibles: {wb.sheetnames}")
    ws = wb[sheet_name]
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)

    headers = {}
    for cell in ws[HEADER_ROW]:
        key = normalize(cell.value)
        if key:
            headers[key] = cell.column

    aliases = {
        "NIVEL": ["NIVEL"],
        "CARGO EN ESTRUCTURA": ["CARGO EN ESTRUCTURA"],
        "DENOMINACION GENERICA DEL PUESTO": ["DENOMINACION GENERICA DEL PUESTO"],
        "ESCOLARIDAD REQUERIDA": ["ESCOLARIDAD REQUERIDA"],
        "EXPERIENCIA (AÑOS)": ["EXPERIENCIA (ANOS)"],
        "CONOCIMIENTOS": ["CONOCIMIENTOS"],
        "DESCRIPCION DE PUESTOS (GENERICA)": ["DESCRIPCION DE PUESTOS (GENERICA)"],
    }

    colmap = {}
    for output_name, candidates in aliases.items():
        found = None
        for cand in candidates:
            if normalize(cand) in headers:
                found = headers[normalize(cand)]
                break
        if found is None:
            raise ValueError(f"No se encontró la columna: {output_name}")
        colmap[output_name] = found

    records = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        rec = {k: clean(ws.cell(r, c).value) for k, c in colmap.items()}
        if not any(rec.values()):
            continue
        # Una fila se considera puesto si al menos trae NIVEL o CARGO.
        if not rec["NIVEL"] and not rec["CARGO EN ESTRUCTURA"]:
            continue
        records.append(rec)
    wb.close()
    return records


def json_text(value):
    """Convierte listas JSON en párrafos sin perder sus elementos."""
    if isinstance(value, list):
        return "\n".join(clean(item) for item in value if clean(item))
    return clean(value)


def load_records_json(json_path):
    data = json.loads(json_path.read_text(encoding="utf-8"))
    source_records = data.get("registros") if isinstance(data, dict) else data
    if not isinstance(source_records, list):
        raise ValueError("El JSON debe contener una lista o un campo 'registros'.")

    field_map = {
        "NIVEL": "nivel",
        "CARGO EN ESTRUCTURA": "puesto",
        "DENOMINACION GENERICA DEL PUESTO": "denominacion_generica_puesto",
        "ESCOLARIDAD REQUERIDA": "escolaridad_requerida",
        "EXPERIENCIA (AÑOS)": "experiencia_anios",
        "CONOCIMIENTOS": "conocimientos",
        "DESCRIPCION DE PUESTOS (GENERICA)": "funciones_genericas",
    }
    records = []
    for source in source_records:
        if not isinstance(source, dict):
            continue
        record = {
            output_name: json_text(source.get(json_name))
            for output_name, json_name in field_map.items()
        }
        if record["NIVEL"] or record["CARGO EN ESTRUCTURA"]:
            records.append(record)
    return records


def load_records(source_path, sheet_name=DEFAULT_SHEET):
    suffix = source_path.suffix.lower()
    if suffix == ".json":
        return load_records_json(source_path)
    if suffix == ".xlsx":
        return load_records_excel(source_path, sheet_name)
    raise ValueError("La fuente debe ser un archivo .json o .xlsx")


def build_catalog(records, output_path, logo_path, limit=None):
    doc = Document()
    setup_doc(doc)
    selected = records if limit is None else records[:limit]

    for i, rec in enumerate(selected):
        sec = doc.sections[0] if i == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        set_record_header(sec, rec, logo_path)
        add_requirements(doc, rec)
        add_long_section(doc, "PERFIL DEL PUESTO", "PERFIL DEL PUESTO", rec["CONOCIMIENTOS"])
        add_long_section(
            doc,
            "FUNCIONES",
            "FUNCIONES",
            rec["DESCRIPCION DE PUESTOS (GENERICA)"],
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return len(selected)


def main():
    parser = argparse.ArgumentParser(description="Genera el catálogo de puestos desde JSON o Excel")
    parser.add_argument(
        "fuente",
        nargs="?",
        default="puestos_unificados.json",
        help="Archivo JSON unificado o Excel de entrada",
    )
    parser.add_argument("--salida", default="catalogo_puestos.docx", help="Archivo Word de salida")
    parser.add_argument("--logo", default="ssh_logo_recortado.png", help="Logo de Servicios de Salud de Hidalgo")
    parser.add_argument("--hoja", default=DEFAULT_SHEET, help="Nombre de la hoja del Excel")
    parser.add_argument("--limite", type=int, default=None, help="Generar solo N puestos para prueba")
    args = parser.parse_args()

    records = load_records(Path(args.fuente), args.hoja)
    n = build_catalog(records, Path(args.salida), Path(args.logo), args.limite)
    print(f"Registros detectados: {len(records)}")
    print(f"Fichas generadas: {n}")
    print(f"Salida: {Path(args.salida).resolve()}")


if __name__ == "__main__":
    main()
